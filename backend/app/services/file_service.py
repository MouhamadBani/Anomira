from __future__ import annotations

import os
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Dict, List, Tuple
from urllib.parse import parse_qs, unquote, urlparse
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl import load_workbook

from app.utils.serialization import dataframe_preview

ALLOWED_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".json", ".jsonl"}
DEFAULT_MAX_UPLOAD_MB = "4" if os.getenv("VERCEL") else "8192"
MAX_UPLOAD_SIZE_MB = int(os.getenv("ANOMIRA_MAX_UPLOAD_MB", DEFAULT_MAX_UPLOAD_MB))
MAX_UPLOAD_SIZE_BYTES = MAX_UPLOAD_SIZE_MB * 1024 * 1024 if MAX_UPLOAD_SIZE_MB > 0 else None
MAX_CLOUD_DOWNLOAD_SIZE_MB = int(os.getenv("ANOMIRA_MAX_CLOUD_DOWNLOAD_MB", "8192"))
MAX_CLOUD_DOWNLOAD_SIZE_BYTES = MAX_CLOUD_DOWNLOAD_SIZE_MB * 1024 * 1024 if MAX_CLOUD_DOWNLOAD_SIZE_MB > 0 else None
CLOUD_DOWNLOAD_TIMEOUT_SEC = int(os.getenv("ANOMIRA_CLOUD_DOWNLOAD_TIMEOUT_SEC", "1800"))
LARGE_DATASET_THRESHOLD_ROWS = 120000
STREAM_CHUNK_SIZE = 8 * 1024 * 1024
TB_MODE_TRIGGER_MB = int(os.getenv("ANOMIRA_TB_MODE_TRIGGER_MB", "512"))
CHUNKED_TB_FORMATS = [".csv", ".tsv", ".txt", ".jsonl"]


class FileValidationError(Exception):
    pass


class FileParsingError(Exception):
    pass


def get_upload_capabilities() -> Dict[str, Any]:
    return {
        "max_upload_size_mb": MAX_UPLOAD_SIZE_MB if MAX_UPLOAD_SIZE_MB > 0 else None,
        "max_upload_size_bytes": MAX_UPLOAD_SIZE_BYTES,
        "max_cloud_download_size_mb": MAX_CLOUD_DOWNLOAD_SIZE_MB if MAX_CLOUD_DOWNLOAD_SIZE_MB > 0 else None,
        "max_cloud_download_size_bytes": MAX_CLOUD_DOWNLOAD_SIZE_BYTES,
        "local_upload_unlimited": MAX_UPLOAD_SIZE_MB <= 0,
        "cloud_download_unlimited": MAX_CLOUD_DOWNLOAD_SIZE_MB <= 0,
        "supported_formats": sorted(ALLOWED_EXTENSIONS),
        "supported_cloud_url_protocols": ["https", "http"],
        "cloud_url_upload_enabled": True,
        "large_dataset_threshold_rows": LARGE_DATASET_THRESHOLD_ROWS,
        "tb_ready_async_job_enabled": True,
        "chunked_tb_mode_formats": CHUNKED_TB_FORMATS,
        "tb_mode_trigger_mb": TB_MODE_TRIGGER_MB,
        "notes": (
            "Full-data analysis is supported. Streamed temp-file ingestion is used for large datasets, "
            "cloud URL uploads are supported, TB-ready async job mode is available for very large text datasets, "
            "and GIS coordinate checks are included. "
            "On Vercel, direct browser uploads are capped by Vercel's function payload limit; use Cloud URL analysis "
            "for larger files. "
            "Set ANOMIRA_MAX_UPLOAD_MB<=0 and/or ANOMIRA_MAX_CLOUD_DOWNLOAD_MB<=0 for unbounded-mode limits."
        ),
    }


def validate_filename(filename: str) -> Tuple[str, str]:
    if not filename:
        raise FileValidationError("A filename is required")

    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise FileValidationError(f"Unsupported file type '{suffix}'. Allowed: {allowed}")

    return filename, suffix


def _create_temp_dataset_path(suffix: str) -> Path:
    fd, path = tempfile.mkstemp(prefix="anomira_", suffix=suffix)
    os.close(fd)
    return Path(path)


def validate_file_size(file_bytes: bytes, *, max_size_bytes: int | None = MAX_UPLOAD_SIZE_BYTES, label: str = "File") -> None:
    if max_size_bytes is None:
        return
    file_size = len(file_bytes)
    if file_size > max_size_bytes:
        raise FileValidationError(
            f"{label} size {round(file_size / (1024 * 1024), 2)}MB exceeds max allowed size of "
            f"{round(max_size_bytes / (1024 * 1024), 2)}MB."
        )


def _validate_file_path_size(file_path: str | Path, *, max_size_bytes: int | None, label: str) -> None:
    if max_size_bytes is None:
        return
    file_size = Path(file_path).stat().st_size
    if file_size > max_size_bytes:
        raise FileValidationError(
            f"{label} size {round(file_size / (1024 * 1024), 2)}MB exceeds max allowed size of "
            f"{round(max_size_bytes / (1024 * 1024), 2)}MB."
        )


def load_dataframe(file_bytes: bytes, filename: str) -> pd.DataFrame:
    _, suffix = validate_filename(filename)
    if not file_bytes:
        raise FileParsingError("Uploaded file is empty")
    validate_file_size(file_bytes)

    try:
        if suffix == ".csv":
            df = pd.read_csv(BytesIO(file_bytes), low_memory=False)
        elif suffix in {".tsv", ".txt"}:
            df = pd.read_csv(BytesIO(file_bytes), sep="\t", low_memory=False)
        elif suffix == ".xlsx":
            df = pd.read_excel(BytesIO(file_bytes), engine="openpyxl")
        elif suffix == ".jsonl":
            df = pd.read_json(BytesIO(file_bytes), lines=True)
        else:
            df = pd.read_json(BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001
        raise FileParsingError(f"Failed to parse {filename}: {exc}") from exc

    if df.empty:
        raise FileParsingError("Dataset contains no rows")

    df.columns = [str(col).strip() for col in df.columns]
    return df


def load_dataframe_from_path(
    file_path: str | Path,
    filename: str,
    *,
    max_size_bytes: int | None = None,
) -> pd.DataFrame:
    _, suffix = validate_filename(filename)
    path = str(file_path)
    if max_size_bytes is not None:
        _validate_file_path_size(path, max_size_bytes=max_size_bytes, label="File")

    try:
        if suffix == ".csv":
            df = pd.read_csv(path, low_memory=False)
        elif suffix in {".tsv", ".txt"}:
            df = pd.read_csv(path, sep="\t", low_memory=False)
        elif suffix == ".xlsx":
            df = pd.read_excel(path, engine="openpyxl")
        elif suffix == ".jsonl":
            df = pd.read_json(path, lines=True)
        else:
            df = pd.read_json(path)
    except Exception as exc:  # noqa: BLE001
        raise FileParsingError(f"Failed to parse {filename}: {exc}") from exc

    if df.empty:
        raise FileParsingError("Dataset contains no rows")

    df.columns = [str(col).strip() for col in df.columns]
    return df


def _read_csv_preview_from_path(file_path: str | Path, sep: str) -> pd.DataFrame:
    return pd.read_csv(str(file_path), sep=sep, nrows=10, low_memory=False)


def _count_text_rows_from_path(file_path: str | Path) -> int:
    line_count = 0
    with open(file_path, "rb") as source:
        while True:
            chunk = source.read(STREAM_CHUNK_SIZE)
            if not chunk:
                break
            line_count += chunk.count(b"\n")
    return max(0, int(line_count) - 1)


def _xlsx_metadata_from_path(file_path: str | Path) -> Tuple[int, List[str], pd.DataFrame]:
    workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        max_row = worksheet.max_row or 0
        header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), tuple())
        columns = [str(col).strip() for col in header_values if col is not None]
    finally:
        workbook.close()

    preview_df = pd.read_excel(str(file_path), engine="openpyxl", nrows=10)
    preview_df.columns = [str(col).strip() for col in preview_df.columns]
    row_count = max(0, int(max_row) - 1)
    return row_count, columns or [str(col) for col in preview_df.columns], preview_df


def extract_upload_metadata_from_path(
    file_path: str | Path,
    filename: str,
    *,
    max_size_bytes: int | None = MAX_UPLOAD_SIZE_BYTES,
    size_label: str = "File",
) -> Dict[str, Any]:
    _, suffix = validate_filename(filename)
    path = str(file_path)
    _validate_file_path_size(path, max_size_bytes=max_size_bytes, label=size_label)

    try:
        if suffix == ".csv":
            preview_df = _read_csv_preview_from_path(path, sep=",")
            row_count = _count_text_rows_from_path(path)
            columns = [str(col).strip() for col in preview_df.columns]
        elif suffix in {".tsv", ".txt"}:
            preview_df = _read_csv_preview_from_path(path, sep="\t")
            row_count = _count_text_rows_from_path(path)
            columns = [str(col).strip() for col in preview_df.columns]
        elif suffix == ".xlsx":
            row_count, columns, preview_df = _xlsx_metadata_from_path(path)
        elif suffix == ".jsonl":
            preview_df = pd.read_json(path, lines=True, nrows=10)
            with open(path, "rb") as source:
                row_count = int(sum(1 for _ in source))
            columns = [str(col).strip() for col in preview_df.columns]
        else:
            full_df = pd.read_json(path)
            if full_df.empty:
                raise FileParsingError("Dataset contains no rows")
            full_df.columns = [str(col).strip() for col in full_df.columns]
            return {
                "row_count": int(len(full_df)),
                "column_count": int(len(full_df.columns)),
                "columns": [str(col) for col in full_df.columns],
                "preview": dataframe_preview(full_df, limit=10),
            }
    except FileParsingError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise FileParsingError(f"Failed to parse {filename}: {exc}") from exc

    if preview_df.empty:
        raise FileParsingError("Dataset contains no rows")

    preview_df.columns = [str(col).strip() for col in preview_df.columns]
    preview = dataframe_preview(preview_df, limit=10)
    column_names = [str(col).strip() for col in columns] if columns else [str(col) for col in preview_df.columns]
    row_count = max(int(row_count), int(len(preview_df)))

    return {
        "row_count": int(row_count),
        "column_count": int(len(column_names)),
        "columns": column_names,
        "preview": preview,
    }


def extract_upload_metadata(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    _, suffix = validate_filename(filename)
    if not file_bytes:
        raise FileParsingError("Uploaded file is empty")
    validate_file_size(file_bytes)

    temp_path = _create_temp_dataset_path(suffix)
    try:
        temp_path.write_bytes(file_bytes)
        return extract_upload_metadata_from_path(temp_path, filename)
    finally:
        temp_path.unlink(missing_ok=True)


def _resolve_cloud_filename(url: str, filename_hint: str | None = None) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise FileValidationError("Only HTTP/HTTPS cloud URLs are supported.")

    if filename_hint:
        filename = filename_hint.strip()
    else:
        filename = unquote(Path(parsed.path).name or "").strip()

    if filename:
        try:
            validate_filename(filename)
            return filename
        except FileValidationError:
            filename = ""

    query = parse_qs(parsed.query)
    query_candidates: list[str] = []
    for key, values in query.items():
        key_l = key.lower()
        if key_l in {"filename", "file", "name", "key", "path", "download"}:
            query_candidates.extend(values)
    if not query_candidates:
        for values in query.values():
            query_candidates.extend(values)

    for raw in query_candidates:
        candidate = unquote(Path(raw).name or "").strip()
        if not candidate:
            continue
        try:
            validate_filename(candidate)
            return candidate
        except FileValidationError:
            continue

    if not filename:
        raise FileValidationError("Could not infer filename from URL. Provide a filename with a supported extension.")

    validate_filename(filename)
    return filename


def download_cloud_url_to_temp_file(
    url: str,
    filename_hint: str | None = None,
    progress_callback: Callable[[int, int | None], None] | None = None,
) -> tuple[str, Path, int]:
    filename = _resolve_cloud_filename(url, filename_hint=filename_hint)
    _, suffix = validate_filename(filename)
    temp_path = _create_temp_dataset_path(suffix)

    try:
        request = Request(url, headers={"User-Agent": "Anomira/1.0"})
        downloaded_size = 0
        with urlopen(request, timeout=CLOUD_DOWNLOAD_TIMEOUT_SEC) as response, temp_path.open("wb") as destination:
            content_length = response.headers.get("Content-Length")
            declared_size: int | None = None
            if content_length:
                declared_size = int(content_length)
                if MAX_CLOUD_DOWNLOAD_SIZE_BYTES is not None and declared_size > MAX_CLOUD_DOWNLOAD_SIZE_BYTES:
                    raise FileValidationError(
                        f"Remote file is too large ({round(declared_size / (1024 * 1024), 2)}MB). "
                        f"Max cloud download size is {MAX_CLOUD_DOWNLOAD_SIZE_MB}MB."
                    )
            if progress_callback:
                progress_callback(downloaded_size, declared_size)

            while True:
                chunk = response.read(STREAM_CHUNK_SIZE)
                if not chunk:
                    break
                downloaded_size += len(chunk)
                if MAX_CLOUD_DOWNLOAD_SIZE_BYTES is not None and downloaded_size > MAX_CLOUD_DOWNLOAD_SIZE_BYTES:
                    raise FileValidationError(
                        f"Remote file exceeded max cloud download size of {MAX_CLOUD_DOWNLOAD_SIZE_MB}MB."
                    )
                destination.write(chunk)
                if progress_callback:
                    progress_callback(downloaded_size, declared_size)

        if downloaded_size <= 0:
            raise FileParsingError("Cloud URL resolved to an empty dataset.")

        return filename, temp_path, downloaded_size
    except FileValidationError:
        temp_path.unlink(missing_ok=True)
        raise
    except Exception as exc:  # noqa: BLE001
        temp_path.unlink(missing_ok=True)
        raise FileParsingError(f"Failed to download from cloud URL: {exc}") from exc
